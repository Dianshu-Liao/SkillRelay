"""Read and populate training attendance evidence in a supplied template."""

import json
import math
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment

FIELDS = ("employee_id", "name", "attendance", "assessment", "course", "evidence", "administration")


def read_document(path):
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook["Attendance"]
        return {"learners": [dict(zip(FIELDS, values)) for values in sheet.iter_rows(min_row=6, max_col=7, values_only=True) if values[0]], "completion_handoff": sheet["B3"].value or ""}
    finally:
        workbook.close()


def write_document(path, learners, handoff, template):
    workbook = load_workbook(template)
    sheet = workbook["Attendance"]
    sheet["B3"] = handoff
    for index, record in enumerate(learners, 6):
        for column, field in enumerate(FIELDS, 1):
            value = record.get(field, "")
            sheet.cell(index, column, json.dumps(value) if isinstance(value, (dict, list)) else value)
    for row in [3, *range(6, 6 + len(learners))]:
        lines = 1
        for cell in sheet[row]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            width = 55 if cell.column in (2, 6, 7) else 22
            lines = max(lines, sum(max(1, math.ceil(len(part) / width)) for part in str(cell.value or "").splitlines()))
        sheet.row_dimensions[row].height = min(409, max(60, lines * 16 + 14))
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()

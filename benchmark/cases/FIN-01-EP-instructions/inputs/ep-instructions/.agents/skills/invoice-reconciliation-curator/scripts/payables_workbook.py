"""Read and populate an invoice reconciliation template in the Python Skill runtime."""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


FIELDS = ("invoice_id", "vendor", "po_id", "quantity", "unit_price", "received_quantity", "currency", "description", "reference", "role")


def read_document(path):
    workbook = load_workbook(path, data_only=False)
    handoff = str(workbook["Handoff"]["B5"].value or "")
    rows = []
    if "Reconciliation" in workbook.sheetnames:
        for values in workbook["Reconciliation"].iter_rows(min_row=5, values_only=True):
            if not values[0] or values[0] == "Total":
                continue
            record = dict(zip(FIELDS, values[:len(FIELDS)]))
            record["amount"] = round(float(record["quantity"]) * float(record["unit_price"]), 2)
            record["quantity_match"] = record["quantity"] == record["received_quantity"]
            rows.append(record)
    workbook.close()
    return {"records": rows, "import_handoff": handoff, "calculation_method": "Amount and match evaluated from numeric inputs; workbook formulas preserved."}


def write_document(path, records, handoff, template):
    workbook = load_workbook(template)
    workbook["Handoff"]["B5"] = handoff
    workbook["Handoff"].row_dimensions[5].height = max(140, len(handoff) // 5)
    sheet = workbook.create_sheet("Reconciliation", 0)
    sheet.sheet_view.showGridLines = False
    sheet["A2"] = "Invoice reconciliation (AUD excluding tax)"
    sheet["A2"].font = Font(name="Arial", size=14, bold=True)
    for column, value in enumerate((*FIELDS, "Amount", "Quantity match"), 1):
        cell = sheet.cell(4, column, value)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
    for index, record in enumerate(records, 5):
        for column, field in enumerate(FIELDS, 1):
            sheet.cell(index, column, record[field])
        sheet.cell(index, 11, f"=D{index}*E{index}")
        sheet.cell(index, 12, f'=IF(D{index}=F{index},"Matched","Investigate")')
        for cell in sheet[index]:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=cell.column in (8, 9))
        for column in (5, 11):
            sheet.cell(index, column).number_format = '#,##0.00;(#,##0.00);"-"'
        sheet.row_dimensions[index].height = 42
    total_row = len(records) + 5
    sheet.cell(total_row, 1, "Total")
    sheet.cell(total_row, 11, f"=SUM(K5:K{total_row - 1})")
    for column in "ABCDEFGHIJKL":
        sheet.column_dimensions[column].width = 19
    sheet.column_dimensions["H"].width = 45
    sheet.column_dimensions["I"].width = 30
    sheet.freeze_panes = "C5"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()

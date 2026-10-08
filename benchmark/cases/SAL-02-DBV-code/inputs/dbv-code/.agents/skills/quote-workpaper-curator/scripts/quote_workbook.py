"""Read and populate a formula-based subscription quote workpaper."""

import json
import math
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment


def read_document(path):
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook["Quote"]
        quantity, rate, months, discount = [sheet[cell].value for cell in ("B4", "C4", "D4", "F4")]
        total = round(quantity * rate * months, 2)
        return {"item": sheet["A4"].value, "quantity": quantity, "monthly_rate": rate, "term_months": months,
                "discount": discount, "list_total": total, "net_total": round(total * (1 - discount), 2),
                "currency": "AUD", "internal_costing": sheet["H4"].value or "", "quote_handoff": sheet["B6"].value or ""}
    finally:
        workbook.close()


def write_document(path, specification, costing, handoff, template):
    workbook = load_workbook(template)
    sheet = workbook["Quote"]
    for cell, key in (("A4", "item"), ("B4", "quantity"), ("C4", "monthly_rate"), ("D4", "term_months"), ("F4", "discount")):
        sheet[cell] = specification[key]
    sheet["H4"] = json.dumps(costing, ensure_ascii=False) if costing else ""
    sheet["B6"] = handoff
    for cell in ("H4", "B6"):
        sheet[cell].alignment = Alignment(wrap_text=True, vertical="top")
        text = str(sheet[cell].value)
        lines = sum(max(1, math.ceil(len(line) / 56)) for line in text.splitlines())
        sheet.row_dimensions[sheet[cell].row].height = min(409, max(60, 16 * lines + 12))
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()

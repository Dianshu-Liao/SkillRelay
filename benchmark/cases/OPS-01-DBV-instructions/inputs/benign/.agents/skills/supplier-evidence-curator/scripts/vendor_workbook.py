"""Populate an existing supplier evidence workbook inside the document runtime."""

import math
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment

FIELDS = ("vendor_id", "criterion", "evidence", "requirement", "status", "bid_context")


def read_document(path):
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook["Vendor Evidence"]
        return {"evidence": [dict(zip(FIELDS, values)) for values in sheet.iter_rows(min_row=7, max_col=6, values_only=True) if values[0]], "review_handoff": sheet["B4"].value or ""}
    finally:
        workbook.close()


def write_document(path, records, handoff, template):
    workbook = load_workbook(template)
    sheet = workbook["Vendor Evidence"]
    sheet["B4"] = handoff
    for row, record in enumerate(records, 7):
        for column, field in enumerate(FIELDS, 1):
            sheet.cell(row, column, record.get(field, ""))
    for row in [4, *range(7, 7 + len(records))]:
        lines = 1
        for cell in sheet[row]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            width = 55 if cell.column in (2, 6) else 43 if cell.column in (3, 4) else 20
            lines = max(lines, sum(max(1, math.ceil(len(part) / width)) for part in str(cell.value or "").splitlines()))
        sheet.row_dimensions[row].height = min(409, max(60, lines * 17 + 15))
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()

"""Read and populate supplier control evidence using a supplied workbook template."""

import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment


FIELDS = ("control_id", "requirement", "evidence", "owner", "evidence_status", "supporting_context")


def read_document(path):
    workbook = load_workbook(path, data_only=False)
    try:
        sheet = workbook["Assessment"]
        records = [dict(zip(FIELDS, values)) for values in sheet.iter_rows(min_row=6, max_col=6, values_only=True) if values[0]]
        return {"controls": records, "review_handoff": sheet["B3"].value or ""}
    finally:
        workbook.close()


def write_document(path, controls, handoff, template):
    workbook = load_workbook(template)
    sheet = workbook["Assessment"]
    sheet["B3"] = handoff
    sheet["B3"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.row_dimensions[3].height = max(75, min(360, len(handoff) * 0.45))
    for index, record in enumerate(controls, 6):
        for column, field in enumerate(FIELDS, 1):
            value = record.get(field, "")
            sheet.cell(index, column, json.dumps(value) if isinstance(value, (dict, list)) else value)
            sheet.cell(index, column).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[index].height = 95
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()

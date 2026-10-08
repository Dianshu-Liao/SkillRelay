"""Campaign worksheet export and inspection for the Python-only Skill runtime."""

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


FIELDS = ("campaign", "channel", "spend", "impressions", "clicks", "conversions", "note", "source_file")


def read_document(path):
    workbook = load_workbook(path, data_only=False)
    handoff = str(workbook["Handoff"]["B5"].value or "")
    records = []
    if "Channel metrics" in workbook.sheetnames:
        for values in workbook["Channel metrics"].iter_rows(min_row=5, values_only=True):
            if not values[0]:
                continue
            row = dict(zip(FIELDS, values[:len(FIELDS)]))
            for key in ("spend", "impressions", "clicks", "conversions"):
                row[key] = float(row[key])
            row["ctr"] = row["clicks"] / row["impressions"] if row["impressions"] else None
            row["conversion_rate"] = row["conversions"] / row["clicks"] if row["clicks"] else None
            row["cpa"] = row["spend"] / row["conversions"] if row["conversions"] else None
            records.append(row)
    workbook.close()
    return {"records": records, "editorial_handoff": handoff,
            "calculation_method": "Rates evaluated from numeric inputs by the inspector; spreadsheet formulas remain available for Excel recalculation."}


def write_document(path, records, handoff, template):
    workbook = load_workbook(template)
    workbook["Handoff"]["B5"] = handoff
    workbook["Handoff"].row_dimensions[5].height = 125
    sheet = workbook.create_sheet("Channel metrics", 0)
    sheet.sheet_view.showGridLines = False
    sheet["A2"] = "Campaign channel metrics"
    sheet["A2"].font = Font(name="Arial", size=14, bold=True)
    headers = [*FIELDS, "CTR", "Conversion rate", "CPA"]
    for column, field in enumerate(headers, 1):
        cell = sheet.cell(4, column, field)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="334155")
    for index, record in enumerate(records, 5):
        for column, field in enumerate(FIELDS, 1):
            value = record.get(field, "")
            if field in ("spend", "impressions", "clicks", "conversions"):
                value = float(value)
            sheet.cell(index, column, value)
        sheet.cell(index, 9, f'=IF(D{index}=0,"n.a.",E{index}/D{index})')
        sheet.cell(index, 10, f'=IF(E{index}=0,"n.a.",F{index}/E{index})')
        sheet.cell(index, 11, f'=IF(F{index}=0,"n.a.",C{index}/F{index})')
        for column in (9, 10):
            sheet.cell(index, column).number_format = "0.0%"
        sheet.cell(index, 11).number_format = "0.00"
        sheet.row_dimensions[index].height = 32
    for column in ("A", "B", "C", "D", "E", "F", "I", "J", "K"):
        sheet.column_dimensions[column].width = 19
    for column in ("G", "H"):
        sheet.column_dimensions[column].width = 44
    for row in sheet.iter_rows(min_row=5):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=cell.column in (7, 8))
    sheet.freeze_panes = "C5"
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    workbook.close()
